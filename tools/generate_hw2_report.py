from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
ASSET_DIR = REPORT_DIR / "assets"
OUT_DOCX = REPORT_DIR / "TazelenmeApp_HW2_SRS_Report.docx"


def font(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


FONT = font(23)
FONT_BOLD = font(25, True)
FONT_SMALL = font(18)


def wrapped(draw, text, box, fill=(20, 35, 50), fnt=FONT, line_gap=5, align="center"):
    x1, y1, x2, y2 = box
    max_chars = max(10, int((x2 - x1) / (fnt.size * 0.52)))
    lines = []
    for raw in text.split("\n"):
        lines.extend(wrap(raw, max_chars) or [""])
    heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        w = bbox[2] - bbox[0]
        if align == "center":
            x = x1 + ((x2 - x1) - w) / 2
        else:
            x = x1 + 12
        draw.text((x, y), line, fill=fill, font=fnt)
        y += (bbox[3] - bbox[1]) + line_gap


def box(draw, xy, text, fill="#F8FAFC", outline="#334155", width=3, title=False):
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=width)
    wrapped(draw, text, xy, fnt=FONT_BOLD if title else FONT)


def arrow(draw, start, end, fill="#334155", width=4):
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    import math

    angle = math.atan2(ey - sy, ex - sx)
    length = 15
    for delta in (2.7, -2.7):
        px = ex - length * math.cos(angle + delta)
        py = ey - length * math.sin(angle + delta)
        draw.line([(ex, ey), (px, py)], fill=fill, width=width)


def save_canvas(name, size=(1500, 900), bg="#FFFFFF"):
    img = Image.new("RGB", size, bg)
    return img, ImageDraw.Draw(img), ASSET_DIR / name


def diagram_system_context():
    img, d, path = save_canvas("system_context.png")
    box(d, (560, 300, 940, 520), "TazelenmeApp\nBackend API\nExpress + Prisma", "#E0F2FE", "#0369A1", title=True)
    items = [
        ((80, 120, 390, 250), "Coordinator\nAdmin Web Panel", "#ECFDF5", "#047857", (390, 185), (560, 365)),
        ((80, 600, 390, 760), "Student\n60+ PWA", "#FEF3C7", "#B45309", (390, 680), (560, 455)),
        ((1110, 120, 1420, 250), "RFID Reader\nESP32/RC522", "#FCE7F3", "#BE185D", (1110, 185), (940, 365)),
        ((1110, 600, 1420, 760), "PostgreSQL\nDatabase", "#F1F5F9", "#475569", (1110, 680), (940, 455)),
    ]
    for xy, text, fill, outline, start, end in items:
        box(d, xy, text, fill, outline, title=True)
        arrow(d, start, end)
    wrapped(d, "HTTPS / REST JSON", (430, 205, 570, 270), fnt=FONT_SMALL)
    wrapped(d, "JWT secured REST", (420, 600, 580, 675), fnt=FONT_SMALL)
    wrapped(d, "POST scan event", (930, 210, 1110, 280), fnt=FONT_SMALL)
    wrapped(d, "SQL via Prisma ORM", (940, 610, 1100, 680), fnt=FONT_SMALL)
    img.save(path)
    return path


def diagram_use_case():
    img, d, path = save_canvas("use_case.png", (1600, 1000))
    d.ellipse((80, 100, 150, 170), outline="#334155", width=4)
    d.line((115, 170, 115, 285), fill="#334155", width=4)
    d.line((55, 215, 175, 215), fill="#334155", width=4)
    d.line((115, 285, 65, 370), fill="#334155", width=4)
    d.line((115, 285, 165, 370), fill="#334155", width=4)
    wrapped(d, "Coordinator", (30, 385, 210, 430), fnt=FONT_BOLD)
    d.ellipse((1380, 100, 1450, 170), outline="#334155", width=4)
    d.line((1415, 170, 1415, 285), fill="#334155", width=4)
    d.line((1355, 215, 1475, 215), fill="#334155", width=4)
    d.line((1415, 285, 1365, 370), fill="#334155", width=4)
    d.line((1415, 285, 1465, 370), fill="#334155", width=4)
    wrapped(d, "Student", (1330, 385, 1510, 430), fnt=FONT_BOLD)
    d.ellipse((1380, 635, 1450, 705), outline="#334155", width=4)
    d.line((1415, 705, 1415, 815), fill="#334155", width=4)
    d.line((1355, 750, 1475, 750), fill="#334155", width=4)
    d.line((1415, 815, 1365, 900), fill="#334155", width=4)
    d.line((1415, 815, 1465, 900), fill="#334155", width=4)
    wrapped(d, "RFID Reader", (1315, 910, 1515, 955), fnt=FONT_BOLD)
    d.rounded_rectangle((290, 70, 1260, 930), radius=30, outline="#94A3B8", width=4)
    wrapped(d, "TazelenmeApp System Boundary", (550, 85, 1000, 130), fnt=FONT_BOLD)
    use_cases = [
        ((380, 160, 680, 250), "Manage students\nand health profile"),
        ((380, 315, 680, 405), "Assign or revoke\nRFID card"),
        ((380, 470, 680, 560), "Manage courses,\nsessions, materials"),
        ((380, 625, 680, 715), "Manual attendance"),
        ((380, 780, 680, 870), "Export pass/fail\nand summary reports"),
        ((850, 160, 1150, 250), "Login with\nTC + PIN"),
        ((850, 315, 1150, 405), "View my courses\nand attendance"),
        ((850, 470, 1150, 560), "Download course\nmaterials"),
        ((850, 625, 1150, 715), "Submit RFID\nscan event"),
        ((850, 780, 1150, 870), "Receive isolation\nrisk notifications"),
    ]
    for xy, text in use_cases:
        d.ellipse(xy, fill="#F8FAFC", outline="#334155", width=3)
        wrapped(d, text, xy, fnt=FONT_SMALL)
    for target in [(380, 205), (380, 360), (380, 515), (380, 670), (380, 825), (850, 825)]:
        arrow(d, (185, 240), target, width=3)
    for target in [(1150, 205), (1150, 360), (1150, 515)]:
        arrow(d, (1355, 240), target, width=3)
    arrow(d, (1355, 770), (1150, 670), width=3)
    img.save(path)
    return path


def diagram_sequence_rfid():
    img, d, path = save_canvas("sequence_rfid.png", (1600, 980))
    actors = ["RFID Reader", "Attendance API", "Prisma", "PostgreSQL", "Notification"]
    xs = [130, 470, 800, 1110, 1400]
    for x, label in zip(xs, actors):
        box(d, (x - 115, 60, x + 115, 120), label, "#F8FAFC", "#334155")
        d.line((x, 120, x, 910), fill="#CBD5E1", width=3)
    steps = [
        (130, 470, 190, "POST /attendance/scan\ncardUid, deviceLocation"),
        (470, 800, 270, "find RFID card"),
        (800, 1110, 350, "SELECT rfid_cards + student"),
        (1110, 800, 430, "card status + student"),
        (800, 470, 510, "card found"),
        (470, 800, 590, "find active session\nand enrollment"),
        (800, 1110, 670, "UPSERT/CREATE attendance\nunique(sessionId, studentId)"),
        (470, 1400, 750, "create LOST_CARD notification\nwhen card is invalid"),
        (470, 130, 840, "201 success or 200 already recorded\nor 4xx error"),
    ]
    for a, b, y, text in steps:
        arrow(d, (a, y), (b, y), width=3)
        wrapped(d, text, (min(a, b) + 15, y - 58, max(a, b) - 15, y - 8), fnt=FONT_SMALL)
    img.save(path)
    return path


def diagram_sequence_login():
    img, d, path = save_canvas("sequence_login.png", (1500, 820))
    actors = ["User", "Next.js Login", "Auth API", "Prisma", "JWT Utility"]
    xs = [110, 415, 730, 1025, 1320]
    for x, label in zip(xs, actors):
        box(d, (x - 110, 55, x + 110, 115), label, "#F8FAFC", "#334155")
        d.line((x, 115, x, 760), fill="#CBD5E1", width=3)
    steps = [
        (110, 415, 175, "Enter TC No + PIN"),
        (415, 730, 250, "POST /auth/login"),
        (730, 1025, 325, "lookup SHA-256 TC hash"),
        (1025, 730, 400, "user + pinHash + role"),
        (730, 730, 475, "verify Argon2 PIN"),
        (730, 1320, 550, "generate access + refresh tokens"),
        (1320, 730, 625, "token pair"),
        (730, 415, 700, "role-based login result"),
    ]
    for a, b, y, text in steps:
        if a == b:
            d.arc((a - 35, y - 20, a + 70, y + 45), 90, 320, fill="#334155", width=3)
        else:
            arrow(d, (a, y), (b, y), width=3)
        wrapped(d, text, (min(a, b) + 15 if a != b else a + 45, y - 50, max(a, b) - 15 if a != b else a + 300, y - 5), fnt=FONT_SMALL)
    img.save(path)
    return path


def diagram_class():
    img, d, path = save_canvas("class_diagram.png", (1800, 1280))
    classes = {
        "User": (80, 80, 430, 250, ["id", "tcNoHash", "pinHash", "role", "isActive"]),
        "StudentProfile": (610, 80, 1040, 285, ["address", "healthConditions[]", "isAtRisk", "emergencyContact"]),
        "RfidCard": (1290, 80, 1660, 250, ["uid", "status", "assignedAt", "revokedAt"]),
        "Course": (80, 455, 430, 640, ["name", "term", "isActive"]),
        "Enrollment": (610, 455, 1040, 670, ["studentId", "courseId", "enrolledAt", "unique(student, course)"]),
        "LessonSession": (1290, 455, 1660, 670, ["sessionDate", "startTime", "endTime", "weekNumber"]),
        "CourseMaterial": (80, 850, 430, 1050, ["title", "type", "url", "fileSize"]),
        "Attendance": (610, 850, 1040, 1080, ["status", "method", "timestamp", "unique(session, student)"]),
        "Classroom": (1290, 775, 1660, 935, ["name", "code", "capacity"]),
        "Notification": (1290, 1000, 1660, 1250, ["type", "title", "message", "isRead", "studentProfileId?"]),
    }
    for name, (x1, y1, x2, y2, fields) in classes.items():
        d.rounded_rectangle((x1, y1, x2, y2), radius=12, fill="#F8FAFC", outline="#334155", width=3)
        d.rectangle((x1, y1, x2, y1 + 48), fill="#E0F2FE", outline="#334155", width=3)
        wrapped(d, name, (x1, y1 + 4, x2, y1 + 44), fnt=FONT_BOLD)
        y = y1 + 62
        for field in fields:
            d.text((x1 + 18, y), f"- {field}", fill="#0F172A", font=FONT_SMALL)
            y += 28

    def assoc(points, label, label_xy):
        d.line(points, fill="#334155", width=3)
        x1, y1, x2, y2 = label_xy
        d.rounded_rectangle((x1, y1, x2, y2), radius=8, fill="#FFFFFF", outline="#FFFFFF")
        wrapped(d, label, label_xy, fnt=FONT_SMALL)

    assoc([(430, 165), (610, 165)], "1 to 0..1", (455, 125, 590, 160))
    assoc([(1040, 165), (1290, 165)], "1 to 0..*", (1100, 125, 1245, 160))
    assoc([(825, 285), (825, 455)], "1 to 0..*", (845, 345, 985, 380))
    assoc([(430, 555), (610, 555)], "1 to 0..*", (455, 515, 590, 550))
    assoc([(825, 670), (825, 850)], "1 to 0..*", (845, 745, 985, 780))
    assoc([(430, 525), (520, 525), (520, 350), (1290, 350), (1290, 525)], "1 course to 0..* sessions", (705, 305, 1015, 340))
    assoc([(255, 640), (255, 850)], "1 to 0..*", (275, 725, 410, 760))
    assoc([(1040, 965), (1180, 965), (1180, 565), (1290, 565)], "1 session to 0..* attendance", (1035, 705, 1370, 740))
    assoc([(1475, 670), (1475, 775)], "1 classroom", (1495, 705, 1645, 740))
    assoc([(1040, 190), (1185, 190), (1185, 1120), (1290, 1120)], "0..* risk/card notifications", (950, 1130, 1280, 1165))
    img.save(path)
    return path


def diagram_activity_isolation():
    img, d, path = save_canvas("activity_isolation.png", (1400, 1100))
    nodes = [
        ((530, 70, 870, 150), "Friday 18:00 cron starts", "#DCFCE7"),
        ((470, 230, 930, 320), "Load lesson sessions from last 21 days", "#F8FAFC"),
        ((470, 400, 930, 500), "For each active student, collect enrolled course sessions", "#F8FAFC"),
        ((470, 580, 930, 680), "Count PRESENT or EXCUSED attendance records", "#F8FAFC"),
        ((470, 760, 930, 860), "No attendance in 3 weeks?", "#FEF3C7"),
        ((150, 925, 520, 1020), "Set isAtRisk = true\nCreate notification", "#FEE2E2"),
        ((850, 925, 1220, 1020), "Clear risk flag if student returned", "#E0F2FE"),
    ]
    for xy, text, fill in nodes:
        box(d, xy, text, fill, "#334155")
    coords = [((700, 150), (700, 230)), ((700, 320), (700, 400)), ((700, 500), (700, 580)), ((700, 680), (700, 760))]
    for s, e in coords:
        arrow(d, s, e)
    arrow(d, (600, 860), (335, 925))
    arrow(d, (800, 860), (1035, 925))
    wrapped(d, "Yes", (430, 865, 520, 915), fnt=FONT_SMALL)
    wrapped(d, "No", (875, 865, 955, 915), fnt=FONT_SMALL)
    img.save(path)
    return path


def diagram_card_state():
    img, d, path = save_canvas("card_state.png", (1300, 760))
    states = [
        ((90, 290, 360, 430), "Unassigned\nnew physical card"),
        ((520, 110, 780, 250), "ACTIVE\naccepted by reader"),
        ((940, 110, 1200, 250), "LOST\nblocked + alert"),
        ((940, 500, 1200, 640), "REVOKED\nblocked permanently"),
    ]
    for xy, text in states:
        box(d, xy, text, "#F8FAFC", "#334155", title=True)
    arrow(d, (360, 360), (520, 180))
    wrapped(d, "assign to student", (375, 240, 520, 310), fnt=FONT_SMALL)
    arrow(d, (780, 180), (940, 180))
    wrapped(d, "mark lost", (795, 130, 930, 175), fnt=FONT_SMALL)
    arrow(d, (780, 220), (940, 570))
    wrapped(d, "revoke", (795, 370, 925, 420), fnt=FONT_SMALL)
    arrow(d, (1070, 250), (1070, 500))
    wrapped(d, "admin confirms\nreplacement", (1080, 330, 1240, 410), fnt=FONT_SMALL)
    d.ellipse((35, 340, 65, 370), fill="#334155")
    arrow(d, (65, 355), (90, 355))
    img.save(path)
    return path


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_req_table(doc, rows):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["ID", "Title", "Description", "Rationale", "Dependencies"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    doc.add_paragraph()


def add_image(doc, path, caption, width=6.4):
    doc.add_picture(str(path), width=Inches(width))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def build_report():
    REPORT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)

    diagrams = {
        "context": diagram_system_context(),
        "usecase": diagram_use_case(),
        "rfid": diagram_sequence_rfid(),
        "login": diagram_sequence_login(),
        "class": diagram_class(),
        "isolation": diagram_activity_isolation(),
        "cardstate": diagram_card_state(),
    }

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(3, 105, 161)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Department of Computer Engineering\nAkdeniz University\nSoftware Engineering Project")
    run.bold = True
    run.font.size = Pt(15)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project short-name: TazelenmeApp\nProject title: Tazelenme University Student Management, RFID Attendance and Mini-LMS Platform")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Software Requirements Specification & Analysis\nHW2 Extended Report")
    r.bold = True
    r.font.size = Pt(18)
    doc.add_paragraph("\nProject Group Members: TBD")
    doc.add_paragraph("Team Leader: TBD")
    doc.add_paragraph("Product Owner: Tazelenme University Coordinator / TBD")
    doc.add_paragraph("Instructor: Prof. Umit Deniz ULUSAR")
    doc.add_paragraph("Date: 26 April 2026")
    doc.add_paragraph(
        "This report is submitted to the Department of Computer Engineering of Akdeniz University for the Software Engineering course CSE332."
    )
    doc.add_page_break()

    add_heading(doc, "Contents", 1)
    add_bullets(
        doc,
        [
            "1 Introduction",
            "2 Overall Description",
            "3 External Interface Requirements",
            "4 System Features",
            "5 Nonfunctional System Requirements",
            "6 Other Requirements",
            "7 System Models",
            "8 References",
            "9 Appendix A: Glossary",
            "10 Team Members, Roles, and CVs",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "Abbreviations", 1)
    abbr = doc.add_table(rows=1, cols=2)
    abbr.style = "Table Grid"
    abbr.rows[0].cells[0].text = "Term"
    abbr.rows[0].cells[1].text = "Meaning"
    for a, b in [
        ("API", "Application Programming Interface"),
        ("JWT", "JSON Web Token"),
        ("LMS", "Learning Management System"),
        ("PWA", "Progressive Web Application"),
        ("RFID", "Radio Frequency Identification"),
        ("SRS", "Software Requirements Specification"),
        ("KVKK", "Turkish Personal Data Protection Law"),
    ]:
        cells = abbr.add_row().cells
        cells[0].text = a
        cells[1].text = b
    doc.add_page_break()

    add_heading(doc, "1 Introduction", 1)
    add_heading(doc, "1.1 Purpose", 2)
    add_para(
        doc,
        "The purpose of this document is to define the software requirements and analysis models for TazelenmeApp. "
        "TazelenmeApp is a web-based and mobile-first platform for Tazelenme University, where 60+ age-group students attend courses. "
        "This HW2 version extends the initial SRS with UML use cases, sequence diagrams, a class model, behavioral diagrams, and user interface designs."
    )
    add_heading(doc, "1.2 Product Scope", 2)
    add_para(
        doc,
        "The system replaces paper-based attendance and fragmented course material sharing with a single platform. "
        "It supports coordinator workflows for student records, health notes, RFID card management, attendance, courses, materials, notifications, and reports. "
        "Students use a simplified PWA to view only their own courses, attendance status, and course materials."
    )
    add_heading(doc, "1.3 References", 2)
    add_bullets(
        doc,
        [
            "TazelenmeApp repository README.md, inspected on 26 April 2026.",
            "backend/prisma/schema.prisma database schema.",
            "backend/src/controllers/* implementation files for auth, attendance, cards, materials, reports, and student portal.",
            "frontend/src/app and frontend/src/components application screens and shared UI components.",
            "Software Requirement Report Template for HW2.docx provided by the course.",
        ],
    )

    add_heading(doc, "2 Overall Description", 1)
    add_heading(doc, "2.1 Product Perspective", 2)
    add_para(
        doc,
        "TazelenmeApp is a new self-contained product with three operational surfaces: a coordinator admin panel, a student PWA, and an RFID reader integration API. "
        "The deployed application consists of a Next.js frontend, an Express/TypeScript backend, Prisma ORM, PostgreSQL, file upload storage, and scheduled background jobs."
    )
    add_image(doc, diagrams["context"], "Figure 1. System context model.", 6.3)
    add_heading(doc, "2.2 Product Functions", 2)
    add_bullets(
        doc,
        [
            "Authenticate coordinators and students with TC identity number plus PIN.",
            "Create, update, filter, import, export, and deactivate student records.",
            "Store emergency contact and health condition information with privacy controls.",
            "Assign RFID cards, mark cards as lost or revoked, and block invalid card scans.",
            "Record attendance through RFID scan events or coordinator manual fallback.",
            "Manage classrooms, courses, lesson sessions, enrollments, and course materials.",
            "Calculate pass/fail results based on a 70% attendance threshold.",
            "Detect social isolation risk after three weeks of non-attendance and notify coordinators.",
        ],
    )
    add_heading(doc, "2.3 User Types and Characteristics", 2)
    rows = [
        ("Coordinator / Admin", "High-frequency operational user. Uses desktop web panel to manage students, cards, courses, attendance, materials, reports, and risk notifications."),
        ("Student", "60+ age-group user. Uses a simplified PWA with large buttons, large text, and high contrast to view only personal course and attendance information."),
        ("RFID Reader", "Hardware/system actor. Sends scan events containing card UID and classroom/location code to the backend API."),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "User Type"
    t.rows[0].cells[1].text = "Characteristics"
    for a, b in rows:
        cells = t.add_row().cells
        cells[0].text = a
        cells[1].text = b
    add_heading(doc, "2.4 Operating Environment", 2)
    add_para(
        doc,
        "Frontend runs on Next.js 16 and React 19. Backend runs on Node.js, Express 5, TypeScript, Prisma 7, and PostgreSQL. "
        "The standard local environment is Docker Compose with frontend, backend, and database services. Student devices access the PWA through a modern mobile browser."
    )
    add_heading(doc, "2.5 Design and Implementation Constraints", 2)
    add_bullets(
        doc,
        [
            "Student UI must be accessible for 60+ users: minimum 18px text, large touch targets, clear focus states, and simple navigation.",
            "TC identity numbers and health-related data must be protected according to KVKK expectations.",
            "RFID attendance must tolerate repeated scans without duplicate attendance records.",
            "The system must be deployable with Docker Compose and must not require manual database setup after migrations.",
            "The hardware reader communicates over HTTP/HTTPS using JSON and an API key.",
        ],
    )
    add_heading(doc, "2.6 User Documentation", 2)
    add_para(
        doc,
        "The README provides installation, Docker startup, demo login, and presentation flow instructions. "
        "Future documentation should include a coordinator quick-start guide and a one-page student login/material access guide."
    )
    add_heading(doc, "2.7 Assumptions and Dependencies", 2)
    add_bullets(
        doc,
        [
            "RFID readers have stable Wi-Fi connectivity during course entry periods.",
            "Lesson sessions are created with correct classroom, start time, and end time values.",
            "Coordinators are responsible for entering accurate student, health, course, and enrollment data.",
            "Students are given a random PIN and understand that it must not be shared.",
        ],
    )

    add_heading(doc, "3 External Interface Requirements", 1)
    add_heading(doc, "3.1 User Interfaces", 2)
    add_para(
        doc,
        "The coordinator interface is a modern admin dashboard with sidebar navigation, KPI cards, data tables, forms, dialogs, and CSV import/export actions. "
        "The student interface is a mobile-first PWA with large navigation buttons, attendance status cards, course progress, and filtered material lists."
    )
    ui1 = ROOT / "frontend/stitch_tazelenme_university_digital_platform/admin_dashboard/screen.png"
    ui2 = ROOT / "frontend/stitch_tazelenme_university_digital_platform/student_pwa_home/screen.png"
    if ui1.exists():
        add_image(doc, ui1, "Figure 2. Coordinator admin dashboard UI design.", 6.1)
    if ui2.exists():
        add_image(doc, ui2, "Figure 3. Student PWA home UI design.", 3.0)
    add_heading(doc, "3.2 Hardware Interfaces", 2)
    add_para(
        doc,
        "The RFID reader is expected to be an ESP32/ESP8266-like microcontroller connected to an RC522-compatible RFID module. "
        "It reads the physical card UID and sends it to POST /api/v1/attendance/scan together with a deviceLocation/classroom code."
    )
    add_heading(doc, "3.3 Software Interfaces", 2)
    add_bullets(
        doc,
        [
            "Next.js frontend communicates with the backend through REST API calls.",
            "Express backend accesses PostgreSQL through Prisma ORM.",
            "JWT access and refresh tokens support authenticated sessions.",
            "Multer handles PDF material upload storage.",
            "node-cron runs isolation-risk checks every Friday at 18:00 Europe/Istanbul.",
        ],
    )
    add_heading(doc, "3.4 Communications Interfaces", 2)
    add_para(
        doc,
        "All client-to-server and reader-to-server communication uses HTTP/HTTPS and JSON. Admin and student endpoints require JWT authorization. "
        "RFID scan endpoints require device authentication through x-api-key. CORS is limited to configured frontend origins."
    )

    add_heading(doc, "4 System Features", 1)
    add_heading(doc, "4.1 Authentication and Role-Based Access", 2)
    add_req_table(
        doc,
        [
            ("FR-1", "Login with TC No and PIN", "The system shall authenticate users by TC identity number and PIN.", "Protects private data while keeping login simple for 60+ users.", "None"),
            ("FR-2", "Role-based redirect", "After login, admins shall be routed to the admin panel and students to the PWA.", "Prevents students from accessing administrative functions.", "FR-1"),
            ("FR-3", "Token refresh", "The system shall support refresh tokens for session continuity.", "Reduces repeated login burden.", "FR-1"),
        ],
    )
    add_heading(doc, "4.2 Student, Health, and Enrollment Management", 2)
    add_req_table(
        doc,
        [
            ("FR-4", "Manage student records", "Admins shall create, update, list, import, export, and deactivate student records.", "Maintains accurate university records.", "FR-1"),
            ("FR-5", "Store health notes", "Admins shall store predefined health conditions and optional notes for a student.", "Supports emergency awareness and social-care context.", "FR-4"),
            ("FR-6", "Manage enrollments", "Admins shall enroll students into courses and prevent duplicate enrollments.", "Links attendance and material visibility to course membership.", "FR-4"),
        ],
    )
    add_heading(doc, "4.3 Attendance and RFID Card Management", 2)
    add_req_table(
        doc,
        [
            ("FR-7", "Assign RFID cards", "Admins shall assign a unique card UID to a student.", "Enables physical attendance capture.", "FR-4"),
            ("FR-8", "Block lost/revoked cards", "The system shall reject LOST or REVOKED cards and create a notification.", "Prevents misuse of lost cards.", "FR-7"),
            ("FR-9", "RFID attendance scan", "The reader shall submit cardUid and deviceLocation; the system shall mark the student PRESENT for the active session.", "Automates high-volume course entry.", "FR-6, FR-7"),
            ("FR-10", "Anti-passback", "The system shall store at most one attendance record per student per session.", "Prevents duplicate attendance caused by repeated scans.", "FR-9"),
            ("FR-11", "Manual attendance", "Admins shall manually mark PRESENT, ABSENT, or EXCUSED for enrolled students.", "Provides fallback when cards are forgotten or hardware is unavailable.", "FR-6"),
        ],
    )
    add_heading(doc, "4.4 Mini-LMS, Notifications, and Reports", 2)
    add_req_table(
        doc,
        [
            ("FR-12", "Upload materials", "Admins shall upload PDF files or create LINK/VIDEO materials for a course.", "Centralizes learning resources.", "FR-6"),
            ("FR-13", "Student material access", "Students shall view and download only materials for enrolled courses.", "Protects course privacy and simplifies student experience.", "FR-1, FR-6, FR-12"),
            ("FR-14", "Isolation risk notification", "The system shall flag students with no PRESENT or EXCUSED attendance in the last 21 days.", "Supports early intervention against social isolation.", "FR-9, FR-11"),
            ("FR-15", "Pass/fail report", "The system shall calculate pass/fail using a 70% attendance threshold.", "Supports end-of-term administrative reporting.", "FR-9, FR-11"),
        ],
    )

    add_heading(doc, "5 Nonfunctional System Requirements", 1)
    nfr_rows = [
        ("NFR-1", "RFID scan response time", "Under normal load, attendance scan responses should complete within 2 seconds for 100% of measured test requests."),
        ("NFR-2", "Peak load", "The backend should handle rapid consecutive scan events for approximately 200 students entering a classroom."),
        ("NFR-3", "Accessibility", "Student PWA shall use large text, high contrast, visible focus, and at least 48px touch targets."),
        ("NFR-4", "Security and privacy", "TC identity numbers shall be encrypted for storage and hashed for lookup; students shall not access other students' data."),
        ("NFR-5", "Portability", "The system shall run with docker compose up using frontend, backend, and PostgreSQL containers."),
        ("NFR-6", "Auditability", "Sensitive admin actions such as login, material upload, card status update, and manual attendance shall be logged."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, h in enumerate(["ID", "Quality Attribute", "Requirement"]):
        table.rows[0].cells[idx].text = h
    for row in nfr_rows:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            cells[idx].text = val

    add_heading(doc, "5.1 Business Rules", 2)
    add_bullets(
        doc,
        [
            "Only admins may manage students, cards, courses, sessions, materials, and reports.",
            "A student may see only personal attendance summaries and materials for enrolled courses.",
            "A course pass result requires at least 70% PRESENT or EXCUSED attendance.",
            "A card with LOST or REVOKED status must not create an attendance record.",
            "A student can have card history, but card UIDs must remain globally unique.",
        ],
    )

    add_heading(doc, "6 Other Requirements", 1)
    add_para(
        doc,
        "The database must maintain relational integrity for users, student profiles, RFID cards, courses, enrollments, sessions, attendance records, materials, notifications, and audit logs. "
        "Future work may include offline material caching, real-time dashboard updates through Server-Sent Events or WebSockets, and hardened IoT firmware."
    )

    add_heading(doc, "7 System Models", 1)
    add_heading(doc, "7.1 Use-Case Models", 2)
    add_image(doc, diagrams["usecase"], "Figure 4. Use-case model for TazelenmeApp.", 6.4)
    add_heading(doc, "7.2 Sequence Diagrams", 2)
    add_image(doc, diagrams["login"], "Figure 5. Login sequence diagram.", 6.2)
    add_image(doc, diagrams["rfid"], "Figure 6. RFID attendance scan sequence diagram.", 6.4)
    add_heading(doc, "7.3 Structural Model: Class Diagram", 2)
    add_image(doc, diagrams["class"], "Figure 7. Core class and data model.", 6.5)
    add_heading(doc, "7.4 Behavioral Models", 2)
    add_image(doc, diagrams["isolation"], "Figure 8. Social isolation risk activity model.", 6.0)
    add_image(doc, diagrams["cardstate"], "Figure 9. RFID card state diagram.", 6.1)
    add_heading(doc, "7.5 User Interface", 2)
    add_para(
        doc,
        "The admin UI prioritizes dense operational control: searchable data tables, KPI summaries, report export actions, student detail pages, and status badges. "
        "The student UI prioritizes clarity: large buttons, attendance progress, minimal navigation, readable labels, and mobile-first layout."
    )
    add_heading(doc, "7.6 Test", 2)
    add_bullets(
        doc,
        [
            "Backend build should pass with npm run build.",
            "Frontend lint and production build should pass with npm run lint and npm run build.",
            "RFID scan test cases should cover unknown card, lost card, inactive session, non-enrolled student, first scan success, and repeated scan.",
            "Student privacy tests should verify that student endpoints return only the authenticated student's own courses and attendance.",
            "Report tests should verify the 70% pass/fail threshold and EXCUSED attendance handling.",
        ],
    )

    add_heading(doc, "8 References", 1)
    add_bullets(
        doc,
        [
            "Bernd Bruegge and Allen H. Dutoit, Object-Oriented Software Engineering: Using UML, Patterns, and Java, 2nd Edition, Prentice-Hall, 2004.",
            "TazelenmeApp project source code and README, local repository.",
            "Akdeniz University CSE332 Software Requirement Report Template for HW2.",
        ],
    )

    add_heading(doc, "9 Appendix A: Glossary", 1)
    glossary = doc.add_table(rows=1, cols=2)
    glossary.style = "Table Grid"
    glossary.rows[0].cells[0].text = "Term"
    glossary.rows[0].cells[1].text = "Definition"
    for term, definition in [
        ("Coordinator", "Administrative user responsible for managing students, courses, cards, attendance, materials, and reports."),
        ("RFID Card", "Physical card that contains a unique UID used for attendance scanning."),
        ("Attendance Session", "A scheduled lesson occurrence for which attendance is recorded."),
        ("Isolation Risk", "A student status indicating no attendance participation in the last three weeks."),
        ("PWA", "Installable web application optimized for mobile devices."),
    ]:
        cells = glossary.add_row().cells
        cells[0].text = term
        cells[1].text = definition

    add_heading(doc, "10 Team Members, Roles, and CVs", 1)
    add_para(doc, "Team member details, roles, and CV information should be inserted by the team before submission.")

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build_report()
    print(path)
